"""AI extraction. The model READS documents and returns structured data with sources. It never converts, sums or compares."""
import datetime as dt, hashlib, io, json, os, threading, time
from schema import Doc, Draft, SupportDoc

MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.8-flash")
# Tried in order when MODEL is overloaded (503), retired (404) or out of its daily free quota (429 PerDay).
# The free tier quota is per model per day, so a longer chain also means more free requests.
FALLBACKS = [m for m in os.environ.get("GEMINI_FALLBACKS", "gemini-3.7-flash,gemini-3.6-flash,gemini-3.5-flash,gemini-3.5-flash-lite,gemini-3.1-flash-lite").split(",") if m and m != MODEL]
PROMPT_VERSION = "v2"
CACHE_DIR = os.path.join(os.path.dirname(__file__), ".cache")

SYSTEM = """You are a procurement data extraction engine. You read one vendor's response to an RFx and return structured data.

Rules, in priority order:
1. NEVER guess. If a value is not written in the document, leave it null. Do not infer prices from other lines.
2. Copy numbers, currencies and units EXACTLY as the vendor wrote them. Do NOT convert currency, units or pack sizes. 'Rs 5,000 per 100 pcs' is price 5000, unit 'per 100 pcs'.
3. Match each vendor line to the RFx line item list provided. Vendors reword and reorder descriptions: match on meaning (product type, ply or cell count, size), not wording. 'Carton 3 ply 12x10x8' is the RFx line '3-ply RSC shipper box 12x10x8 in'. If you still cannot match confidently, set rfx_item_id to null; the buyer will map it.
4. If a statement could apply to several RFx lines (for example '5 ply 42/kg' when both sheets and rolls are priced per kg), emit one line per candidate, set applicability_uncertain true and confidence 0.5 or lower.
5. If a vendor says rates for other items are 'same as last year' or similar, set baseline_reference true and quote the words. Do NOT invent prices for those items.
6. Footnote or paragraph discounts go in 'discounts' with the verbatim sentence. Do not subtract them from prices.
7. For every value, give source_ref and the verbatim source_snippet. For photos, also give box_2d of the price cell.
8. Set confidence below 0.8 if any digit is blurred, smudged, cut off or a unit is unclear. Be honest: a wrong confident number is the worst outcome.
   If a price is partly legible, give your best reading with confidence 0.6 or lower and say in notes which digits are unclear. If it is wholly illegible, emit the line with price null. Never drop a line the vendor quoted.
9. Questionnaire: burst_report is 'Yes' only if a report is provided or attached now. 'Will be furnished later' is 'No'. accepts_60_day_payment is 'No' if the vendor states payment terms shorter than 60 days. Anything unmentioned is 'Not stated'.
10. Terms: record freight wording verbatim. freight_amount_stated is true only if an amount or rate is given.
Return JSON matching the schema. No commentary."""


_CLIENT = None


def _client():
    # One shared client: newer google-genai closes a client's connection when the object is garbage collected,
    # so a throwaway client per call fails with "client has been closed".
    global _CLIENT
    if _CLIENT is None:
        _CLIENT = _new_client()
    return _CLIENT


def _new_client():
    from google import genai
    from google.genai import types
    key = os.environ.get("GEMINI_API_KEY")
    if not key:
        try:
            import streamlit as st
            key = st.secrets.get("GEMINI_API_KEY")
        except Exception:
            key = None
    if not key:
        raise RuntimeError("GEMINI_API_KEY is not set")
    return genai.Client(api_key=key, http_options=types.HttpOptions(timeout=REQUEST_TIMEOUT_S * 1000))


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data + MODEL.encode() + PROMPT_VERSION.encode()).hexdigest()[:24]


def dump_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            cells = [f"{c.coordinate}={c.value!r}" for c in row if c.value not in (None, "")]
            if cells:
                out.append(f"[{ws.title}] " + " | ".join(cells))
    return "\n".join(out)


def dump_docx(data: bytes) -> str:
    from docx import Document
    d = Document(io.BytesIO(data))
    out = [f"P{n}: {p.text}" for n, p in enumerate(d.paragraphs, 1) if p.text.strip()]
    for t, tb in enumerate(d.tables, 1):
        for r, row in enumerate(tb.rows, 1):
            out.append(f"T{t}R{r}: " + " | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def dump_text(data: bytes) -> str:
    return "\n".join(f"L{n}: {l}" for n, l in enumerate(data.decode("utf-8", "ignore").splitlines(), 1) if l.strip())


def kind_of(name: str) -> str:
    n = name.lower()
    for k, exts in {"xlsx": (".xlsx", ".xlsm"), "docx": (".docx",), "text": (".txt", ".eml"), "pdf": (".pdf",),
                    "image": (".jpg", ".jpeg", ".png", ".webp")}.items():
        if n.endswith(exts):
            return k
    return "unknown"


# Fail fast on the free tier: one request may take REQUEST_TIMEOUT_S, a whole question BUDGET_S across the fallback chain.
REQUEST_TIMEOUT_S, BUDGET_S = 45, 90
FALLBACK_SIGNS = ("503", "UNAVAILABLE", "404", "NOT_FOUND", "429", "RESOURCE_EXHAUSTED", "PerDay", "timed out", "Timeout", "DEADLINE_EXCEEDED", "504")
_SKIP = {}  # model -> time until which it is skipped (shared by all sessions: the quota is per project)
_LOCAL = threading.local()


class progress:
    """`with extract.progress(fn):` calls fn(text) as generate() moves through the models, so the screen never looks frozen."""
    def __init__(self, fn):
        self.fn = fn

    def __enter__(self):
        self.prev, _LOCAL.fn = getattr(_LOCAL, "fn", None), self.fn

    def __exit__(self, *exc):
        _LOCAL.fn = self.prev


def _say(text):
    fn = getattr(_LOCAL, "fn", None)
    if fn:
        fn(text)


def _next_quota_reset():
    """Free-tier daily quotas reset at midnight US Pacific time."""
    from zoneinfo import ZoneInfo
    now = dt.datetime.now(ZoneInfo("America/Los_Angeles"))
    return (now + dt.timedelta(days=1)).replace(hour=0, minute=5, second=0, microsecond=0).timestamp()


def _retry(fn, tries=2):
    """Retry only a momentary overload, once and briefly; anything else moves straight on to the next model."""
    for k in range(tries):
        try:
            return fn()
        except Exception as e:
            if k == tries - 1 or not any(s in str(e) for s in ("503", "UNAVAILABLE")):
                raise
            time.sleep(3)


def generate(contents, config, before_attempt=None, budget=None):
    """Call MODEL, then each fallback, skipping models known to be out of quota, retired or just overloaded.
    Returns (response, model_used). before_attempt runs before every try, e.g. to clear a tool trace left by a failed attempt."""
    budget = budget or BUDGET_S
    start, last, tried = time.time(), None, []
    models = [m for m in [MODEL] + FALLBACKS if _SKIP.get(m, 0) < time.time()]
    if not models:
        raise RuntimeError("429 RESOURCE_EXHAUSTED: every model's free allowance is used up for today (PerDay)")
    for n, m in enumerate(models):
        if time.time() - start > budget:
            raise TimeoutError(f"The AI timed out after {budget}s (busy). Tried: {', '.join(tried)}")
        _say(f"Asking the AI ({m})" if not tried else f"{tried[-1]} is busy or out of allowance, trying {m}")

        def call():
            if before_attempt:
                before_attempt()
            return _client().models.generate_content(model=m, contents=contents, config=config)
        try:
            return _retry(call, 2 if n == 0 else 1), m
        except Exception as e:
            msg = str(e)
            if not any(s in msg for s in FALLBACK_SIGNS):
                raise
            if "PerDay" in msg:
                _SKIP[m] = _next_quota_reset()            # out for the day
            elif "404" in msg or "NOT_FOUND" in msg:
                _SKIP[m] = time.time() + 7 * 86400        # retired or unknown model
            else:
                _SKIP[m] = time.time() + 60               # busy or per-minute limit: give it a minute
            tried.append(m)
            last = e
    raise last


def _cache_path(data: bytes, items: list) -> str:
    return os.path.join(CACHE_DIR, f"{file_hash(data + json.dumps([i['id'] for i in items]).encode())}.json")


def is_cached(data: bytes, items: list) -> bool:
    return os.path.exists(_cache_path(data, items))


def extract_response(name: str, data: bytes, items: list, use_cache=True) -> dict:
    """Return a Doc dict for one vendor response file."""
    os.makedirs(CACHE_DIR, exist_ok=True)
    path = _cache_path(data, items)
    if use_cache and os.path.exists(path):
        return json.load(open(path))
    from google.genai import types
    kind = kind_of(name)
    rfx = "RFx line items (id | description | unit):\n" + "\n".join(f"{i['id']} | {i['description'] if 'description' in i else i['desc']} | {i['unit']}" for i in items)
    if kind == "xlsx":
        parts = [types.Part.from_text(text=f"{rfx}\n\nVENDOR FILE '{name}' (Excel, cell addresses shown):\n{dump_xlsx(data)}")]
    elif kind == "docx":
        parts = [types.Part.from_text(text=f"{rfx}\n\nVENDOR FILE '{name}' (Word):\n{dump_docx(data)}")]
    elif kind == "text":
        parts = [types.Part.from_text(text=f"{rfx}\n\nVENDOR EMAIL '{name}':\n{dump_text(data)}")]
    elif kind == "pdf":
        parts = [types.Part.from_text(text=f"{rfx}\n\nThe vendor response is the attached PDF. Use 'page N' in source_ref."),
                 types.Part.from_bytes(data=data, mime_type="application/pdf")]
    elif kind == "image":
        mime = "image/png" if name.lower().endswith(".png") else "image/jpeg"
        parts = [types.Part.from_text(text=f"{rfx}\n\nThe vendor response is the attached photo of a printed rate card, possibly angled or blurred. Use 'row N' (the printed row number) in source_ref and give box_2d for each price."),
                 types.Part.from_bytes(data=data, mime_type=mime)]
    else:
        raise ValueError(f"Unsupported file type: {name}")

    resp, used = generate(parts, types.GenerateContentConfig(system_instruction=SYSTEM, response_mime_type="application/json",
                                                             response_schema=Doc, temperature=0))
    doc = Doc.model_validate_json(resp.text).model_dump()
    ids = {i["id"] for i in items}
    for l in doc["lines"]:
        if l["rfx_item_id"] not in ids:
            l["rfx_item_id"] = None
    doc["_file"] = name
    doc["_kind"] = kind
    doc["_model"] = used
    json.dump(doc, open(path, "w"))
    return doc


SUPPORT_PROMPT_VERSION = "s1"
SUPPORT_SYSTEM = """You read one supporting document attached to a vendor's quote: an ISO certificate, a test report, or something else.
Return structured data. Rules:
1. NEVER guess. Leave a field null if it is not written. Do not judge pass or fail yourself: copy the stated result verbatim.
2. Copy names, numbers and units exactly as written. Dates as YYYY-MM-DD.
3. issued_to is the company the certificate is issued to or the report is about, exactly as written.
4. Put the verbatim lines you relied on in 'evidence'. Set confidence below 0.8 if anything is hard to read.
Return JSON matching the schema. No commentary."""


def extract_support(name: str, data: bytes, use_cache=True) -> dict:
    """Read a certificate or test report into a SupportDoc dict. Code, not the model, decides what it proves."""
    os.makedirs(CACHE_DIR, exist_ok=True)
    path = os.path.join(CACHE_DIR, "sup_" + hashlib.sha256(data + MODEL.encode() + SUPPORT_PROMPT_VERSION.encode()).hexdigest()[:24] + ".json")
    if use_cache and os.path.exists(path):
        return json.load(open(path))
    from google.genai import types
    kind = kind_of(name)
    if kind == "pdf":
        parts = [types.Part.from_text(text=f"Supporting document '{name}' is attached."), types.Part.from_bytes(data=data, mime_type="application/pdf")]
    elif kind == "image":
        parts = [types.Part.from_text(text=f"Supporting document '{name}' is the attached photo."),
                 types.Part.from_bytes(data=data, mime_type="image/png" if name.lower().endswith(".png") else "image/jpeg")]
    elif kind in ("text", "docx"):
        parts = [types.Part.from_text(text=f"Supporting document '{name}':\n" + (dump_docx(data) if kind == "docx" else dump_text(data)))]
    else:
        raise ValueError(f"Unsupported supporting document type: {name}")
    resp, used = generate(parts, types.GenerateContentConfig(system_instruction=SUPPORT_SYSTEM, response_mime_type="application/json",
                                                             response_schema=SupportDoc, temperature=0))
    sup = SupportDoc.model_validate_json(resp.text).model_dump()
    sup.update(_file=name, _kind=kind, _model=used)
    json.dump(sup, open(path, "w"))
    return sup


def support_is_cached(data: bytes) -> bool:
    return os.path.exists(os.path.join(CACHE_DIR, "sup_" + hashlib.sha256(data + MODEL.encode() + SUPPORT_PROMPT_VERSION.encode()).hexdigest()[:24] + ".json"))


def draft_rfx(brief: str) -> dict:
    from google.genai import types
    sys = ("You help a procurement buyer draft an RFx. From the brief, produce line items (description, unit of measure, annual quantity), "
           "5 questionnaire questions (include ISO certification and a test report as quality gates) and commercial terms. "
           "Every line must have an explicit unit of measure so vendors cannot misread it. Do not invent facts the brief does not support; keep quantities plausible.")
    resp, _ = generate(brief, types.GenerateContentConfig(system_instruction=sys, response_mime_type="application/json",
                                                          response_schema=Draft, temperature=0.3))
    return Draft.model_validate_json(resp.text).model_dump()
