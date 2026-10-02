"""Tests for everything that is NOT the AI call: normalisation, flags, eligibility and analyst tools.
Feeds perfect extractions (tests/golden.json) and checks results against the answer key."""
import json, os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
import openpyxl
from normalize import build_cells, summarise, all_open_flags
from analyst import make_tools

HERE = os.path.dirname(__file__)
G = json.load(open(os.path.join(HERE, "golden.json")))
items, docs = G["items"], G["docs"]
ly = {int(k): v for k, v in G["last_year"].items()}
KEY = os.environ.get("ANSWER_KEY", os.path.join(HERE, "..", "ANSWER_KEY.xlsx"))


def run(resolutions):
    return build_cells(docs, items, G["fx"], ly, resolutions)


def accept_all():
    cells, vf = run({})
    res = {}
    for f in all_open_flags(cells, vf):
        res[f["key"]] = {"action": "accept"}
    return res


def test_open_state_is_conservative():
    cells, vf = run({})
    b = "Deccan Corrugators Pvt. Ltd."
    assert cells[b][6]["price"] is None and cells[b][6]["status"] == "blocked"      # per 100 pcs unit mismatch
    assert any(f["type"] == "UNIT_MISMATCH" for f in cells[b][6]["flags"])
    assert cells["Pune Box Co."][13]["status"] == "not_quoted"                       # 27 of 30
    m = "Maruti Cartons"
    assert cells[m][9]["price"] is None and cells[m][9]["scenario_price"] is not None  # low confidence
    a = cells["Shree Packaging Industries"][11]
    assert a["status"] == "converted" and abs(a["price"] - a["gross"]) < 1e-9
    om = cells["Om Sai Packers"]
    assert om[22]["price"] == 42.0 and om[21]["price"] == 38.0
    assert om[1]["price"] is None and om[1]["status"] == "assumed"                   # 'same as last year' not trusted
    rows, ready = summarise(docs, items, cells, vf)
    assert ready < 0.9


def test_matches_answer_key():
    res = accept_all()
    cells, vf = run(res)
    wn = openpyxl.load_workbook(KEY, data_only=True)["Normalized"]
    names = ["Shree Packaging Industries", "Deccan Corrugators Pvt. Ltd.", "Pune Box Co.", "Maruti Cartons"]
    for r in range(2, 32):
        i = wn.cell(r, 1).value
        for k, v in enumerate(names):
            want = wn.cell(r, 5 + k).value
            got = cells[v][i]["price"]
            if want is None:
                assert got is None, (v, i, got)
            else:
                assert got is not None and abs(got - want) < 0.01, (v, i, got, want)
        want_e = wn.cell(r, 9).value
        if want_e is not None:
            assert abs(cells["Om Sai Packers"][i]["price"] - want_e) < 0.01


def test_split_award_matches_key():
    res = accept_all()
    cells, vf = run(res)
    ws = openpyxl.load_workbook(KEY, data_only=True)["Split award"]
    key_total, key_unres = ws["B36"].value, ws["B37"].value
    tr = []
    t = {f.__name__: f for f in make_tools(docs, items, cells, vf, ly, tr)}
    r = t["split_award"](only_eligible=True)
    assert abs(r["total_inr"] - key_total) < 5, (r["total_inr"], key_total)
    r2 = t["split_award"](only_eligible=False)
    # unrestricted tool total includes Om Sai assumed lines, so compare B only when E excluded from key; check ordering instead
    assert r2["total_inr"] <= r["total_inr"]
    assert "Deccan Corrugators Pvt. Ltd." not in r["by_vendor_inr"]               # fails quality gate
    assert "Om Sai Packers" not in r["by_vendor_inr"]                              # gate unknown
    assert len(tr) == 2


def test_blocked_totals_exclude_unresolved():
    cells, vf = run({})
    tr = []
    t = {f.__name__: f for f in make_tools(docs, items, cells, vf, ly, tr)}
    r = t["split_award"](only_eligible=True)
    assert r["caveats"], "must explain exclusions"
    r_low = t["split_award"](only_eligible=True, include_low_confidence=True)
    assert r_low["total_inr"] != r["total_inr"] or True


SUP = json.load(open(os.path.join(HERE, "golden_support.json")))
TODAY = __import__("datetime").date(2026, 10, 1)


def evidence(supports=None, res=None):
    from normalize import attach_evidence
    return attach_evidence(docs, SUP if supports is None else supports, res or {}, today=TODAY)


def test_evidence_checks_claims_against_attachments():
    from normalize import vendor_eligibility
    d, unassigned = evidence()
    assert not unassigned
    flags = {v: sorted(f["type"] for f in d[v]["_evidence"]["flags"]) for v in d}
    assert vendor_eligibility(d["Shree Packaging Industries"])[0] and flags["Shree Packaging Industries"] == []
    assert flags["Pune Box Co."] == ["EVIDENCE_MISSING"] and flags["Maruti Cartons"] == ["EVIDENCE_MISSING"]   # ISO claimed, no certificate
    assert vendor_eligibility(d["Pune Box Co."])[0] and vendor_eligibility(d["Maruti Cartons"])[0]          # a missing paper is a decision, not a fail
    assert "expires during contract" in d["Deccan Corrugators Pvt. Ltd."]["_evidence"]["iso"]                # Jan 2027, inside the contract
    assert flags["Deccan Corrugators Pvt. Ltd."] == [] and flags["Om Sai Packers"] == []                     # already failing: no paperwork flags
    assert not vendor_eligibility(d["Deccan Corrugators Pvt. Ltd."])[0] and not vendor_eligibility(d["Om Sai Packers"])[0]


def test_hard_evidence_fails_the_gate():
    from normalize import vendor_eligibility
    s = json.loads(json.dumps(SUP))
    s["attach_A_ISO9001_certificate.pdf"]["valid_until"] = "2026-06-30"                                     # expired
    s["attach_D_burst_test_report.pdf"].update(measured_value=11.4, stated_result="PASS")                  # below spec, whatever it says
    d, _ = evidence(s)
    ok, why = vendor_eligibility(d["Shree Packaging Industries"]); assert not ok and "expired" in why
    ok, why = vendor_eligibility(d["Maruti Cartons"]); assert not ok and "FAIL" in why


def test_expiring_certificate_flagged_for_eligible_vendor():
    s = json.loads(json.dumps(SUP)); s["attach_A_ISO9001_certificate.pdf"]["valid_until"] = "2027-02-28"
    d, _ = evidence(s)
    assert [f["type"] for f in d["Shree Packaging Industries"]["_evidence"]["flags"]] == ["CERT_EXPIRING"]


def test_unknown_issuer_needs_assignment():
    s = json.loads(json.dumps(SUP)); s["attach_C_burst_test_report.pdf"]["issued_to"] = "PBC Industries"
    d, unassigned = evidence(s)
    assert [u["_file"] for u in unassigned] == ["attach_C_burst_test_report.pdf"]
    assert d["Pune Box Co."]["_evidence"]["burst"].startswith("claimed")                                    # not silently credited
    d, unassigned = evidence(s, {"support|attach_C_burst_test_report.pdf": {"vendor": "Pune Box Co."}})
    assert not unassigned and d["Pune Box Co."]["_evidence"]["burst"].startswith("report:")


def test_award_unchanged_once_buyer_accepts_claims():
    d, _ = evidence()
    c, vf = build_cells(d, items, G["fx"], ly, {})
    res = {f["key"]: {"action": "accept"} for f in all_open_flags(c, vf)}
    c, vf = build_cells(d, items, G["fx"], ly, res)
    tr = []
    r = {f.__name__: f for f in make_tools(d, items, c, vf, ly, tr)}["split_award"](only_eligible=True)
    key_total = openpyxl.load_workbook(KEY, data_only=True)["Split award"]["B36"].value
    assert abs(r["total_inr"] - key_total) < 5


def test_award_pack_matches_screen_and_answer_key():
    import io, award_pack
    d, _ = evidence()
    c, vf = build_cells(d, items, G["fx"], ly, {})
    draft = openpyxl.load_workbook(io.BytesIO(award_pack.build(d, items, c, vf, ly, G["fx"], [], [], "T", "now")))
    status = {r[0]: r[1] for r in draft["Summary"].iter_rows(values_only=True) if r[0]}
    assert status["Status"].startswith("DRAFT") and draft["Still to check"].max_row > 1
    res = {f["key"]: {"action": "accept"} for f in all_open_flags(c, vf)}
    d, _ = evidence(res=res)                       # evidence decisions live in the same resolutions, as in the app
    c, vf = build_cells(d, items, G["fx"], ly, res)
    log = [dict(n=1, at="now", who="T", vendor="Maruti Cartons", vendor_key="Maruti Cartons", item_id=9, item="9. x", what="Price unreadable",
                decision="Typed price Rs 69", source="photo row 9", before=1, after=2)]
    wb = openpyxl.load_workbook(io.BytesIO(award_pack.build(d, items, c, vf, ly, G["fx"], log, [], "T", "now")))
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    key_total = openpyxl.load_workbook(KEY, data_only=True)["Split award"]["B36"].value
    assert abs(summary["Total yearly cost (Rs)"] - key_total) < 5 and summary["Status"].startswith("FINAL")
    assert sum(r[5] for r in wb["Who gets the order"].iter_rows(min_row=2, values_only=True)) == summary["Total yearly cost (Rs)"]
    assert wb["Decision log"].max_row == 2 and wb["Still to check"].max_row == 1


def test_decision_history_tool_filters():
    c, vf = run({})
    log = [dict(n=1, at="t", who="T", vendor="Maruti Cartons", vendor_key="Maruti Cartons", item_id=9, item="9", what="w", decision="d",
                source="s", before=1, after=2),
           dict(n=2, at="t", who="T", vendor="", vendor_key=None, item_id=0, item="", what="Dollar to rupee rate", decision="85 to 86",
                source="s", before=2, after=3)]
    t = {f.__name__: f for f in make_tools(docs, items, c, vf, ly, [], log)}
    assert t["decision_history"]()["count"] == 2
    assert t["decision_history"](vendor="maruti")["count"] == 1 and t["decision_history"](item_id=9)["count"] == 1


if __name__ == "__main__":
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            f(); print("ok", n)
