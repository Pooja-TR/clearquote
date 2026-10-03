import io, os, json, re
import pandas as pd
import streamlit as st
from PIL import Image

import importlib


def _load_helpers():
    """After a git push, Streamlit Cloud re-runs this file but can keep the OLD helper modules in memory, so new names are
    missing (ImportError on the live app). Check each helper for names this version needs; once one is stale, reload it
    and every helper after it, since each imports from the ones before."""
    needs = [("schema", ("SupportDoc",)), ("extract", ("generate", "extract_support", "support_is_cached", "progress")),
             ("normalize", ("attach_evidence", "match_vendor", "EVIDENCE_FLAGS")), ("analyst", ("ask", "EVIDENCE_FLAGS")),
             ("award_pack", ("build",)), ("followup", ("draft", "questions"))]
    versions = {"analyst": 4}
    stale = False
    for name, names in needs:
        try:
            m = importlib.import_module(name)
        except ImportError:  # it imports a name a stale earlier helper lacks; earlier helpers were reloaded above
            stale = True
            continue
        if stale or any(not hasattr(m, n) for n in names) or getattr(m, "API_VERSION", 0) < versions.get(name, 0):
            stale = True
            importlib.reload(m)
    if stale:  # retry anything that failed to import
        for name, _ in needs:
            importlib.import_module(name)


_load_helpers()
import extract
from normalize import build_cells, summarise, all_open_flags, vendor_eligibility, attach_evidence
from analyst import ask, make_tools
import award_pack
import followup

HERE = os.path.dirname(__file__)
SAMPLE = os.path.join(HERE, "sample_data")
MAX_AI_CALLS = 120

st.set_page_config(page_title="ClearQuote", page_icon="📦", layout="wide")
S = st.session_state
for k, v in dict(rfx=[], files={}, docs={}, supp_files={}, supports={}, resolutions={}, turns=[], last_year={}, ai_calls=0, outbox=[],
              log=[], fx_prev=85.0).items():
    S.setdefault(k, v)


def load_sample_rfx():
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(SAMPLE, "00_RFx_line_items_and_questionnaire.xlsx"))["Line items"]
    parts = [x.strip() for x in str(ws["A1"].value or "").split("|")]
    S.rfq = dict(no=parts[0].replace("RFQ No.", "").strip() if parts else "", title=parts[1] if len(parts) > 1 else "Sample request",
                 site=parts[2] if len(parts) > 2 else "")
    S.rfx = [dict(id=r[0], desc=r[1], qty=r[2], unit=r[3]) for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]
    ly = load_workbook(os.path.join(SAMPLE, "00_last_year_awarded_rates.xlsx")).active
    S.last_year = {r[0]: r[4] for r in ly.iter_rows(min_row=2, values_only=True) if r[0]}


def load_sample_responses():
    for f in sorted(os.listdir(SAMPLE)):
        if f[:2] in ("A_", "B_", "C_", "D_", "E_"):
            S.files[f] = open(os.path.join(SAMPLE, f), "rb").read()
        elif f.startswith("attach_"):
            S.supp_files[f] = open(os.path.join(SAMPLE, f), "rb").read()


FLAG_NAMES = {"UNIT_MISMATCH": "Different unit or pack size", "LOW_CONFIDENCE": "AI unsure of this number", "UNREADABLE_PRICE": "Price unreadable",
              "AMBIGUOUS_APPLICABILITY": "Unclear which item", "BASELINE_ASSUMED": "Not re-quoted ('same as last year')",
              "DISCOUNT_FOUND": "Discount offered", "AMBIGUOUS_TERM": "Delivery (freight) cost unknown", "UNMATCHED_LINE": "Couldn't match to an item",
              "CURRENCY_UNKNOWN": "Unknown currency", "EVIDENCE_MISSING": "Said yes, no document", "CERT_EXPIRING": "Certificate expiring",
              "TEST_SCOPE": "Test done on a different box type", "REPORT_OLD": "Old test report"}
TOOL_NAMES = {"split_award": "Who gets the order (split award)", "lowest_price_per_item": "Cheapest vendor per item",
              "vendor_overview": "Vendor summary", "compare_vendors": "Vendor vs vendor", "open_flags": "Things to check",
              "vendor_terms": "Terms and quality documents", "decision_history": "Decision record", "single_vendor_award": "Whole order with one vendor"}
STATUS_NAMES = {"ok": "Confirmed", "converted": "Confirmed, converted from USD", "blocked": "Needs your decision (not in totals)",
                "assumed": "Not re-quoted: last year's price, needs your decision", "not_quoted": "Not quoted"}
GLOSSARY = """
- **Request for quotes (RFx)**: the list of items you ask vendors to price.
- **Who gets the order (award)**: the final decision on which vendor supplies which item. A *split award* gives each item to the cheapest vendor that passed the quality check; a *single award* gives everything to one vendor.
- **Quality check**: must-haves from the request: an ISO 9001 certificate and a burst test report.
- **ISO 9001**: a widely used certificate that a factory runs a proper quality system.
- **Burst test (BF, Mullen)**: how much pressure a box survives before bursting. BF is the board's bursting factor.
- **Delivery (freight)**: transport cost. *Ex-works* means the price excludes delivery; *FOR Pune* means delivered to the plant.
- **Not re-quoted**: the vendor said "same as last year" instead of giving a price.
- **Things to check**: anything the app will not decide for you, such as a smudged price or a price per 100 instead of per box.
"""


COLUMN_WORDS = {"item_id": "Item no.", "inr": "(Rs)", "qty": "Yearly quantity", "lines": "items", "line": "item", "flag": "To check",
                "best": "Cheapest", "runner up": "Next cheapest", "passes quality gate": "Passes quality check", "gate note": "Quality check note",
                "open flags": "Things to check", "of total": "Out of", "iso 9001": "ISO 9001 (said)", "burst report": "Burst report (said)"}


def plain_table(rows):
    """Tool tables for people: readable headings, short vendor names, no raw codes."""
    df = pd.DataFrame(rows)
    if "flag" in df:
        df["flag"] = df["flag"].map(lambda x: FLAG_NAMES.get(x, x))

    def head(c):
        if c in COLUMN_WORDS:
            return COLUMN_WORDS[c]
        h = c.replace("_", " ")
        for k, w in COLUMN_WORDS.items():
            h = re.sub(rf"\b{k}\b", w, h)
        return h[:1].upper() + h[1:]
    df.columns = [short(c) if c in docs else head(c) for c in df.columns]
    return df.map(lambda x: short(x) if isinstance(x, str) and x in docs else x)


def friendly_error(e):
    """AI failures in plain words. Free-tier limits are expected, so say what still works and when the AI is back."""
    msg = str(e)
    if any(k in msg for k in ("RESOURCE_EXHAUSTED", "429", "quota")):
        return ("Today's free AI allowance is used up (this demo runs on Gemini's free tier). Everything already read still works: "
                "the comparison, things to check, decisions and the award pack. Reading new files and answering questions comes back "
                "after midnight US Pacific time (early afternoon in India).")
    if any(k in msg for k in ("503", "UNAVAILABLE", "overloaded", "high demand", "timed out", "Timeout", "DEADLINE_EXCEEDED")):
        return "Google's AI is busy right now. Wait a minute and try again; everything already read still works."
    if "GEMINI_API_KEY" in msg:
        return "The AI key is not set up on this server, so new files and questions can't be read. Everything already read still works."
    return f"The AI step failed: {msg[:200]}"


def plain_args(args):
    """Tool settings in words, for 'How this was worked out'."""
    out = []
    if args.get("only_eligible"):
        out.append("only vendors that passed the quality check")
    elif "only_eligible" in args:
        out.append("all vendors, including those that failed the quality check")
    if args.get("include_low_confidence"):
        out.append("including numbers the AI was unsure of")
    if args.get("vendor_a"):
        out.append(f"{short(args['vendor_a'])} vs {short(args['vendor_b'])}" + (f", items containing '{args['item_filter']}'" if args.get("item_filter") else ""))
    if args.get("vendor"):
        out.append(f"vendor: {short(args['vendor'])}")
    return "; ".join(out)


def pager(total, per_page, key):
    """Previous / next controls. Returns the (start, end) slice for the current page."""
    pages = max(1, -(-total // per_page))
    S[key] = min(max(S.get(key, 1), 1), pages)
    p = S[key]
    if pages > 1:
        c1, c2, c3 = st.columns([1, 3, 1], vertical_alignment="center")
        c1.button("◀ Previous", key=f"{key}_prev", disabled=p == 1, on_click=lambda: S.update({key: p - 1}), width="stretch")
        c2.markdown(f"<div style='text-align:center'>Page {p} of {pages} · showing {(p - 1) * per_page + 1}–{min(p * per_page, total)} of {total}</div>",
                    unsafe_allow_html=True)
        c3.button("Next ▶", key=f"{key}_next", disabled=p == pages, on_click=lambda: S.update({key: p + 1}), width="stretch")
    return (p - 1) * per_page, min(p * per_page, total)


LEGAL_SUFFIXES = r"(,?\s+(pvt\.?|private|ltd\.?|limited|llp|inc\.?|co\.?|company|industries))+\.?$"


def short(v):
    """Display name for a vendor: 'DECCAN CORRUGATORS PVT. LTD.' -> 'Deccan Corrugators'. The full name stays the key everywhere."""
    n = v.title() if v.isupper() else v
    return re.sub(LEGAL_SUFFIXES, "", n, flags=re.I).strip() or n


def fmt_inr(x):
    if x is None:
        return "-"
    if x >= 1e7:
        return f"Rs {x / 1e7:.2f} cr"
    if x >= 1e5:
        return f"Rs {x / 1e5:.2f} lakh"
    return f"Rs {x:,.0f}"


def now_ist():
    return pd.Timestamp.now(tz="Asia/Kolkata").strftime("%d %b %Y, %H:%M IST")


def split_total(res, rate):
    """Split-award total among vendors passing the quality check, for a given set of decisions. Same code as everywhere else."""
    if not (S.docs and S.rfx):
        return None
    d = attach_evidence(S.docs, S.supports, res)[0]
    c, vf = build_cells(d, S.rfx, rate, S.last_year, res)
    return {f.__name__: f for f in make_tools(d, S.rfx, c, vf, S.last_year, [])}["split_award"](only_eligible=True)["total_inr"]


def record(vendor, item_id, what, decision, source, before, after, key=None):
    """Append-only, like a register: nothing is edited or deleted; an undo is a new entry."""
    desc = next((it["desc"] for it in S.rfx if it["id"] == item_id), "")
    S.log.append(dict(n=len(S.log) + 1, at=now_ist(), who=S.get("who") or "Buyer", vendor=short(vendor) if vendor else "",
                      vendor_key=vendor, item_id=item_id, item=f"{item_id}. {desc}" if item_id else ("whole quote" if vendor else ""),
                      what=what, decision=decision, source=source, before=before, after=after, key=key))


def decide(key, res, vendor, item_id, what, decision, source):
    """Save a buyer decision, record it with the total before and after, and refresh."""
    rate = S.get("fx", 85.0)
    before = split_total(S.resolutions, rate)
    S.resolutions = {**S.resolutions, key: res}
    record(vendor, item_id, what, decision, source, before, split_total(S.resolutions, rate), key)
    st.rerun()


def undo(key):
    e = next(e for e in reversed(S.log) if e.get("key") == key and not e["decision"].startswith("Undid"))
    rate = S.get("fx", 85.0)
    before = split_total(S.resolutions, rate)
    S.resolutions = {k: v for k, v in S.resolutions.items() if k != key}
    record(e["vendor_key"], e["item_id"], e["what"], f"Undid #{e['n']} ({e['decision']})", "Buyer reversed an earlier decision",
           before, split_total(S.resolutions, rate), key)


def fx_changed():
    new, old = S.fx, S.fx_prev
    record(None, 0, "Dollar to rupee rate", f"Changed from {old:g} to {new:g}", "Sidebar setting",
           split_total(S.resolutions, old), split_total(S.resolutions, new))
    S.fx_prev = new


with st.sidebar:
    st.header("Settings")
    st.text_input("Your name (for the decision record)", value="Buyer", key="who")
    fx = st.number_input("Dollar to rupee rate (USD to INR)", value=85.0, step=0.5, key="fx", on_change=fx_changed,
                         help="Used to convert prices quoted in US dollars. Applied by the calculations, never by the AI.")
    st.caption("Demo rate. In real use it would be fixed per request, with its date.")
    st.caption(f"AI questions used this session: {S.ai_calls} of {MAX_AI_CALLS}")
    with st.expander("What do these words mean?"):
        st.markdown(GLOSSARY)
    with st.expander("Technical details"):
        st.caption(f"AI assistant: Google Gemini. Main model {extract.MODEL}; if it is busy or out of free quota, a backup model answers. "
                   "Each file and answer shows which model was used.")
        try:
            import subprocess
            ver = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=HERE, capture_output=True, text=True, timeout=5).stdout.strip()
        except Exception:
            ver = ""
        st.caption(f"Version: {ver or 'unknown'}")

# Visual polish only: colours come from .streamlit/config.toml; this adds the record header, cards and tab styling.
st.markdown("""<style>
.block-container {padding-top: 4.2rem;}
.rec {background: #fff; border: 1px solid #dcd8f2; border-left: 5px solid #5a35cf; border-radius: 12px; padding: 14px 20px;
      display: flex; justify-content: space-between; align-items: center; gap: 24px; margin-bottom: 4px;
      box-shadow: 0 2px 10px rgba(28, 26, 74, .05);}
.rec .eyebrow {font-size: .72rem; letter-spacing: .08em; text-transform: uppercase; color: #6b6894; font-weight: 600;}
.rec .title {font-size: 1.25rem; font-weight: 700; color: #1c1a4a; margin: 2px 0;}
.rec .meta {font-size: .86rem; color: #5b5880;}
.rec .meta b {color: #1c1a4a; font-weight: 600;}
.rec .left {flex: 1; min-width: 0;} .rec .right {text-align: right; flex: none;}
.mobile-hint {display: none;}
@media (max-width: 900px) {.rec {flex-wrap: wrap;} .rec .right {text-align: left; flex: 1 1 100%; min-width: 0;}}
@media (max-width: 640px) {
  .block-container {padding-left: 0.8rem; padding-right: 0.8rem;}
  .rec {padding: 12px 14px; gap: 10px;} .rec .title {font-size: 1.1rem;} .rec .big {font-size: 1.3rem;}
  [data-testid="stMetric"] {padding: 6px 12px;} [data-testid="stMetricValue"] {font-size: 1.5rem;}
  .mobile-hint {display: block; margin-top: 4px;}
}
.rec .big {font-size: 1.5rem; font-weight: 700; color: #3d05c6; line-height: 1.2;}
.badge {display: inline-block; font-size: .75rem; font-weight: 700; padding: 2px 10px; border-radius: 999px; margin-left: 8px; vertical-align: middle;}
.badge.draft {background: #fff1e0; color: #b45309;} .badge.final {background: #e3f6ea; color: #11743b;}
.sec {font-size: 1.05rem; font-weight: 700; color: #1c1a4a; margin: 0 0 2px;}
.sub {font-size: .86rem; color: #6b6894; margin: 0 0 10px;}
[data-testid="stTabs"] [role="tablist"] {gap: 6px; margin-top: 6px;}
[data-testid="stTabs"] [role="tab"] {padding: 8px 14px; border-radius: 10px 10px 0 0; font-weight: 600; color: #4a4680;}
[data-testid="stTabs"] [role="tab"][aria-selected="true"] {background: #ebe8fb; color: #3d05c6;}
[data-testid="stTabs"] [role="tab"]:hover {color: #3d05c6;}
[data-testid="stVerticalBlockBorderWrapper"] {background: #fff; box-shadow: 0 2px 10px rgba(28, 26, 74, .05);}
[data-testid="stMetric"] {background: #fff; border: 1px solid #dcd8f2; border-radius: 12px; padding: 12px 16px;
                          box-shadow: 0 2px 10px rgba(28, 26, 74, .05);}
[data-testid="stMetricValue"] {color: #3d05c6; font-weight: 700;}
[data-testid="stDataFrame"] {background: #fff; border-radius: 10px;}
.legend span {display: inline-block; padding: 2px 10px; border-radius: 6px; margin: 0 6px 4px 0; font-size: .82rem; border: 1px solid #dcd8f2;}
</style>""", unsafe_allow_html=True)


def section(title, sub=None):
    st.markdown(f'<div class="sec">{title}</div>' + (f'<div class="sub">{sub}</div>' if sub else ""), unsafe_allow_html=True)


header_slot = st.container()  # filled at the end, once the numbers are known
t1, t2, t3, t_ask, t4, t5 = st.tabs(["1. Request for quotes", "2. Vendor replies", "3. Compare", "4. Ask the AI", "5. Things to check",
                                     "6. Decision record"])

# ---------------------------------------------------------------- tab 1
with t1:
    with st.container(border=True):
        section("What are you buying?", "Describe it in a sentence and the AI drafts the item list, quality questions and terms. "
                                        "Or load the sample: 30 corrugated packaging items for a Pune electronics plant.")
        brief = st.text_area("What do you need to buy?", label_visibility="collapsed", height=80,
                             placeholder="e.g. 30 corrugated packaging items for our Pune plant: shipper boxes, mailers, sheets, pallet caps. 60-day payment.")
        c1, c2, _ = st.columns([1.7, 1.5, 2.4])
        load = c1.button("Load sample request (30 items)", type="primary", width="stretch")
        draft = c2.button("Draft the request with AI", width="stretch", disabled=not brief)
    if load:
        load_sample_rfx(); st.rerun()
    if draft and brief:
        if S.ai_calls >= MAX_AI_CALLS:
            st.error("AI call limit reached for this session.")
        else:
            with st.status("Drafting the request...") as box:
                try:
                    S.ai_calls += 1
                    with extract.progress(lambda t: box.update(label=f"{t}...")):
                        d = extract.draft_rfx(brief)
                    box.update(label="Drafted", state="complete")
                    S.rfx = [dict(id=n, desc=i["description"], unit=i["unit"], qty=i["quantity"]) for n, i in enumerate(d["items"], 1)]
                    S.rfq = dict(no="Draft", title="New request (drafted with AI)", site="")
                    S.draft_extra = d
                except Exception as e:
                    st.error(f"Could not draft the request. {friendly_error(e)}")
    if S.rfx:
        with st.container(border=True):
            section(f"Items in this request ({len(S.rfx)})", "Edit anything before it goes out. Every item states the unit vendors must price in.")
            df = st.data_editor(pd.DataFrame(S.rfx), num_rows="dynamic", width="stretch", hide_index=True, key="items_editor",
                                column_config={"id": st.column_config.NumberColumn("No.", width=60), "desc": st.column_config.TextColumn("Item", width="large"),
                                               "qty": st.column_config.NumberColumn("Yearly quantity", format="localized"),
                                               "unit": st.column_config.TextColumn("Price per (unit)")})
            S.rfx = df.dropna(subset=["desc"]).to_dict("records")
            for n, it in enumerate(S.rfx, 1):
                it["id"] = n if not it.get("id") or pd.isna(it.get("id")) else int(it["id"])
            if "draft_extra" in S:
                with st.expander("Quality questions and terms drafted"):
                    st.write(S.draft_extra["questionnaire"]); st.write(S.draft_extra["terms"])
            if st.button("Send to vendors (simulated)"):
                S.outbox = [f"To: vendor{n}@example.com\nSubject: RFQ\n\nPlease quote on the {len(S.rfx)} items attached. State unit and currency per item." for n in range(1, 6)]
            for m in S.outbox:
                st.code(m, language=None)

# ---------------------------------------------------------------- tab 2
with t2:
    with st.container(border=True):
        section("Vendor replies", "Vendors reply however they like: Excel, PDF, Word, a phone photo, an email. The AI reads each one and "
                                  "notes where every price came from. The samples are 5 such replies plus 5 certificates and test reports.")
        b1, b2, _ = st.columns([1.3, 1.3, 3])
        if b1.button("Load the 5 sample replies", width="stretch"):
            if not S.rfx:
                load_sample_rfx()
            load_sample_responses(); st.rerun()
        read = b2.button("Read replies with AI", type="primary", width="stretch", disabled=not S.files)
        with st.expander("Or upload your own files"):
            u1, u2 = st.columns(2)
            st.caption("This demo runs on a free AI plan: each new file uses one AI request, and each question one to three. "
                       "Try your own files; if the day's allowance runs out, the app says so and everything already read keeps working.")
            up = u1.file_uploader("Quotes", accept_multiple_files=True)
            sup_up = u2.file_uploader("Certificates and test reports", accept_multiple_files=True,
                                      help="Matched to vendors by the company name printed on each document. You can reassign any that do not match.")
        for f in up or []:
            S.files[f.name] = f.getvalue()
        for f in sup_up or []:
            S.supp_files[f.name] = f.getvalue()
    if S.files:
      with st.container(border=True):
        section(f"Quotes received ({len(S.files)})", "'Read by' shows which AI model read each file.")
        st.dataframe(pd.DataFrame([{"File": n, "Type": extract.kind_of(n), "Size": f"{len(b) / 1024:.1f} KB",
                                    "Read": any(d.get("_file") == n for d in S.docs.values()),
                                    "Read by": next((d.get("_model", "") for d in S.docs.values() if d.get("_file") == n), "")} for n, b in S.files.items()]),
                     hide_index=True, width="stretch")
        for m in S.get("read_errors", []):  # kept in session: the page refreshes right after reading
            st.error(m)
        if read:
            S.read_errors = []
            if not S.rfx:
                st.error("Create or load the request first (tab 1).")
            else:
                bar = st.progress(0.0)
                names = list(S.files)
                for name, data in S.supp_files.items():
                    if name in S.supports:
                        continue
                    line = st.empty(); line.write(f"Reading {name} ...")
                    try:
                        if not extract.support_is_cached(data):
                            if S.ai_calls >= MAX_AI_CALLS:
                                raise RuntimeError("AI call limit reached for this session")
                            S.ai_calls += 1
                        with extract.progress(lambda t, n=name: line.write(f"Reading {n}: {t}...")):
                            S.supports[name] = extract.extract_support(name, data)
                    except Exception as e:
                        S.read_errors.append(f"{name}: {friendly_error(e)}")
                for n, name in enumerate(names):
                    if any(d.get("_file") == name for d in S.docs.values()):
                        continue
                    line = st.empty(); line.write(f"Reading {name} ...")
                    try:
                        if not extract.is_cached(S.files[name], S.rfx):
                            if S.ai_calls >= MAX_AI_CALLS:
                                raise RuntimeError("AI call limit reached for this session")
                            S.ai_calls += 1
                        with extract.progress(lambda t, n=name: line.write(f"Reading {n}: {t}...")):
                            doc = extract.extract_response(name, S.files[name], S.rfx)
                        S.docs[doc["vendor_name"]] = doc
                    except Exception as e:
                        S.read_errors.append(f"{name}: {friendly_error(e)}")
                    bar.progress((n + 1) / len(names))
                st.rerun()
    if S.supp_files:
      with st.container(border=True):
        section(f"Certificates and test reports ({len(S.supp_files)})", "Matched to vendors by the company name printed on each document.")
        _, unassigned = attach_evidence(S.docs, S.supports, S.resolutions) if S.docs else ({}, [])
        def owner(n):
            r = S.resolutions.get(f"support|{n}")
            if r:
                return short(r["vendor"]) + " (set by you)"
            from normalize import match_vendor
            sd = S.supports.get(n)
            v = match_vendor(sd["issued_to"], list(S.docs)) if sd and S.docs else None
            return short(v) if v else ("no match" if sd else "")
        def facts(sd):
            if sd["doc_type"] == "iso_certificate":
                return f"valid to {sd.get('valid_until') or 'not shown'}"
            if sd["doc_type"] == "test_report":
                return f"{sd.get('measured_value')} vs min {sd.get('spec_min')}, {sd.get('stated_result') or '?'}"
            return sd.get("title") or ""
        st.dataframe(pd.DataFrame([{"File": n, "Kind": ({"iso_certificate": "ISO certificate", "test_report": "test report"}.get(S.supports[n]["doc_type"], "other") if n in S.supports else "not read yet"),
                                    "Issued to": S.supports[n]["issued_to"] if n in S.supports else "", "Matched vendor": owner(n),
                                    "Key facts": facts(S.supports[n]) if n in S.supports else ""} for n in S.supp_files]),
                     hide_index=True, width="stretch")
        for sd in unassigned:
            with st.container(border=True):
                st.write(f"**{sd['_file']}** is issued to '{sd['issued_to']}', which matches no vendor. Which vendor sent it?")
                pick = st.selectbox("Vendor", list(S.docs), format_func=short, key=f"assign_{sd['_file']}")
                if st.button("Assign", key=f"assignb_{sd['_file']}"):
                    decide(f"support|{sd['_file']}", {"vendor": pick}, pick, 0, "Document matched to vendor",
                           f"Assigned {sd['_file']} to {short(pick)}", f"Document issued to '{sd['issued_to']}'")

def source_text(f):
    """What a decision was based on: the flag's message plus the vendor's own words, for the decision record."""
    c = cells[f["vendor"]][f["item_id"]] if f["item_id"] else None
    src = f" | Source {c['source'].get('ref')}: {c['source'].get('snippet')}" if c and c["source"] else ""
    return f["message"] + src


def show_crop(vendor, c):
    """Show the region of a photo a value was read from (box_2d is 0-1000 [ymin, xmin, ymax, xmax])."""
    fname = docs[vendor].get("_file")
    if not (c["source"] and c["source"].get("box") and fname in S.files and docs[vendor].get("_kind") == "image"):
        return
    try:
        im = Image.open(io.BytesIO(S.files[fname]))
        y0, x0, y1, x1 = c["source"]["box"]
        W, H = im.size
        from PIL import ImageDraw
        im = im.convert("RGB")
        bx = (x0 * W // 1000, y0 * H // 1000, x1 * W // 1000, y1 * H // 1000)
        ImageDraw.Draw(im).rectangle((bx[0] - 6, bx[1] - 6, bx[2] + 6, bx[3] + 6), outline=(230, 120, 0), width=4)
        # Photos are often tilted, so a thin strip can pair a price with the wrong row label. Show a taller band of the
        # full table width with the value boxed, so the buyer can follow the ruled lines to the item name.
        pad = max(70, (bx[3] - bx[1]) * 3)
        crop = im.crop((W // 25, max(0, bx[1] - pad), W - W // 25, min(H, bx[3] + pad)))
        st.image(crop, caption="The value read is boxed in orange. Phone photos are often tilted: follow the table lines to the item, not straight across.", width="stretch")
    except Exception:
        st.caption("Could not crop the photo region.")


# ---------------------------------------------------------------- shared compute
items = S.rfx
docs = attach_evidence(S.docs, S.supports, S.resolutions)[0] if S.docs else {}
cells = vflags = None
split_now = None
if docs and items:
    cells, vflags = build_cells(docs, items, fx, S.last_year, S.resolutions)
    split_now = {f.__name__: f for f in make_tools(docs, items, cells, vflags, S.last_year, [])}["split_award"](only_eligible=True)

# ---------------------------------------------------------------- tab 3
with t3:
    if not cells:
        st.info("Load the request and read the vendor replies first (tabs 1 and 2).")
    else:
        rows, ready = summarise(docs, items, cells, vflags)
        nopen = len(all_open_flags(cells, vflags))
        m1, m3, m4 = st.columns(3)
        m1.metric("Comparison complete", f"{ready:.0%}", help=f"{nopen} things to check (tab 5). Prices waiting on a decision are left out of every total.")
        m3.metric("Pass quality check", f"{sum(r['eligible'] for r in rows)} of {len(rows)}", help="ISO 9001 certified and a burst test report provided.")
        m4.metric("Prices ready to use", f"{sum(c['price'] is not None for v in docs for c in cells[v].values())} of {len(docs) * len(items)}",
                  help="One price per vendor per item. Only confirmed or converted prices count; the rest are left out of totals.")
        st.progress(ready)
        vend = list(docs)
        grid, sty = [], []
        for it in items:
            r, s = {"Item": f"{it['id']}. {it['desc']}", "Unit": it["unit"], "Qty": it["qty"]}, {}
            for v in vend:
                c = cells[v][it["id"]]
                if c["status"] == "ok":
                    r[v], s[v] = f"{c['price']:,.2f}", ""
                elif c["status"] == "converted":
                    r[v], s[v] = f"{c['price']:,.2f}", "background-color:#ebe6ff"
                elif c["status"] in ("blocked", "assumed"):
                    d = c["display_price"]
                    r[v], s[v] = (f"? {d:,.2f}" if d is not None else "? unreadable" if c.get("unreadable") else "? blocked"), "background-color:#ffe9d6"
                else:
                    r[v], s[v] = "not quoted", "background-color:#eeeeee;color:#777"
            r = {short(k) if k in docs else k: x for k, x in r.items()}
            s = {short(k): x for k, x in s.items()}
            grid.append(r); sty.append(s)
        gdf = pd.DataFrame(grid)
        with st.container(border=True):
            section("Prices per item", "Rs per unit asked for in the request. Pick any cell below under 'Where did a price come from?' to see its source."
                                       '<span class="mobile-hint">On a small screen, swipe the table sideways to see each vendor.</span>')
            st.markdown('<div class="legend"><span style="background:#fff">confirmed</span><span style="background:#ebe6ff">converted from USD</span>'
                        '<span style="background:#ffe9d6">? needs your decision, excluded from totals</span>'
                        '<span style="background:#eeeeee;color:#777">not quoted</span></div>', unsafe_allow_html=True)
            st.dataframe(gdf.style.apply(lambda _: pd.DataFrame([{**{"Item": "", "Unit": "", "Qty": ""}, **s} for s in sty], columns=gdf.columns), axis=None),
                         width="stretch", hide_index=True, height=min(35 * (len(items) + 1) + 3, 1100),
                         column_config={"Item": st.column_config.TextColumn(width=240, pinned=True), "Unit": st.column_config.TextColumn(width=86), "Qty": st.column_config.NumberColumn(format="localized")})
        with st.container(border=True):
            section("Vendors", "Totals only add up items with a usable price, so a vendor with fewer usable items is not directly comparable.")
            vdf = pd.DataFrame([{"Vendor": short(r["vendor"]), "Quality check": ("Pass" if r["eligible"] else f"Fail: {r['eligibility_note']}"),
                                 "Items with a usable price": f"{r['lines_usable']} of {len(items)}", "Total for those items": fmt_inr(r["total"]),
                                 "Things to check": r["open_flags"]} for r in rows])
            st.dataframe(vdf.style.map(lambda x: "color:#11743b;font-weight:600" if x == "Pass" else "color:#b42318;font-weight:600", subset=["Quality check"]),
                         hide_index=True, width="stretch")
        with st.container(border=True):
            section("Quality evidence", "What each vendor claimed, and what their attached certificates and test reports actually prove.")
            edf = pd.DataFrame([{"Vendor": short(v), "ISO 9001 claimed": d["questionnaire"]["iso_9001"], "ISO 9001 evidence": d["_evidence"]["iso"],
                                 "Burst report claimed": d["questionnaire"]["burst_report"], "Burst test evidence": d["_evidence"]["burst"],
                                 "Quality check": "Pass" if vendor_eligibility(d)[0] else "Fail"} for v, d in docs.items()])
            weak = lambda x: "color:#b45309;font-weight:600" if ("claimed, no" in str(x) or "during contract" in str(x)) else ("color:#b42318;font-weight:600" if any(w in str(x) for w in ("expired", "FAIL", "not ")) else "")
            st.dataframe(edf.style.map(weak, subset=["ISO 9001 evidence", "Burst test evidence"])
                            .map(lambda x: "color:#11743b;font-weight:600" if x == "Pass" else "color:#b42318;font-weight:600", subset=["Quality check"]),
                         hide_index=True, width="stretch",
                         column_config={"Vendor": st.column_config.TextColumn(width=150), "ISO 9001 claimed": st.column_config.TextColumn("ISO claimed", width=95),
                                        "Burst report claimed": st.column_config.TextColumn("Burst claimed", width=105), "Quality check": st.column_config.TextColumn("Check", width=60)})

        with st.container(border=True):
            section("Where did a price come from?", "Pick an item and a vendor to see the vendor's own words, how the price was worked out, and any photo region.")
            ic1, ic2 = st.columns(2)
            pick = ic1.selectbox("Item", [f"{it['id']}. {it['desc']}" for it in items])
            vsel = ic2.selectbox("Vendor", vend, format_func=short)
            c = cells[vsel][int(pick.split(".")[0])]
            st.write(f"Status: {STATUS_NAMES.get(c['status'], c['status'])}")
            if c["raw"]:
                st.write(f"As quoted: {c['raw']}")
                st.write(f"How it was worked out: {c['receipt']}")
            if c["source"]:
                st.write(f"Source: {c['source'].get('ref')}")
                st.code(c["source"].get("snippet") or "", language=None)
                show_crop(vsel, c)
            for f in c["flags"]:
                st.write(f"To check ({FLAG_NAMES.get(f['type'], f['type'])}): {f['message']}")


# ---------------------------------------------------------------- ask tab
with t_ask:
    if not cells:
        st.info("Load the request and read the vendor replies first (tabs 1 and 2).")
    else:
        section("Ask about this comparison", "Plain questions, plain answers. The AI picks a calculation, the code runs it, "
                                             "and every answer shows how it was worked out and what to watch out for.")
        qs = {"Missing items": "Which vendors did not quote all items, and which items are missing?",
              "Cheapest per item (chart)": "Who is cheapest per item? Show it as a chart.",
              "Unknown delivery costs": "Which quotes have delivery (freight) or other charges we have no number for?",
              "Who gets the order + saving": "If each item goes to the cheapest vendor that passed the quality check, who gets the order, what is the total, and what is the saving versus last year?",
              "One vendor for everything": "If one vendor that passed the quality check gets the whole order, who should it be, what is the total, and what is the saving versus last year?",
              "What is unresolved?": "What is still unresolved, and could any of it change who gets the order?"}
        picked = st.pills("Suggested questions", list(qs), key=f"sugg{len(S.turns)}", help="Click to ask. Hover a finished answer's question to see it in full.")
        with st.form("ask_form", clear_on_submit=True, border=False):
            fc1, fc2 = st.columns([6, 1], vertical_alignment="bottom")
            typed = fc1.text_input("Ask anything about this comparison", placeholder="e.g. Compare Shree and Maruti on 7-ply boxes")
            sent = fc2.form_submit_button("Ask", type="primary", width="stretch")
        question = (typed.strip() if sent and typed.strip() else None) or (qs[picked] if picked else None)
        if question:
            if S.ai_calls >= MAX_AI_CALLS:
                st.error("AI call limit reached for this session.")
            else:
                hist = [x for t in S.turns[-3:] for x in (("user", t["q"]), ("assistant", t["a"]))]
                with st.status("Asking the AI...", expanded=True) as box:
                    shown = st.empty()
                    def on_progress(text):
                        box.update(label=f"{text}...")
                        shown.caption("On the free tier a busy model is skipped and the next one tried; this can take up to a minute and a half.")
                    try:
                        S.ai_calls += 1
                        with extract.progress(on_progress):
                            ans, trace, model = ask(question, hist, docs, items, cells, vflags, S.last_year, S.log)
                        box.update(label=f"Answered by {model}", state="complete")
                    except Exception as e:
                        ans, trace, model = friendly_error(e), [], None
                        box.update(label="The AI could not answer", state="error")
                S.turns.append(dict(n=len(S.turns) + 1, q=question, a=ans, trace=trace, model=model,
                                    at=pd.Timestamp.now(tz="Asia/Kolkata").strftime("%H:%M IST")))
                st.rerun()

        if S.turns:
            hc1, hc2 = st.columns([3, 2], vertical_alignment="bottom")
            hc1.markdown(f"**Answers** ({len(S.turns)}, newest first)")
            needle = hc2.text_input("Filter answers", placeholder="Filter by a word, e.g. freight", label_visibility="collapsed",
                                    key="turn_filter", on_change=lambda: S.update(turn_page=1))
            shown = [t for t in reversed(S.turns) if not needle or needle.lower() in (t["q"] + " " + t["a"]).lower()]
            if not shown:
                st.caption("No answers match that filter.")
            a, b = pager(len(shown), 5, "turn_page")
            for t in shown[a:b]:
                with st.container(border=True):
                    st.markdown(f"**Q{t['n']}. {t['q']}**")
                    st.caption(f"{t['at']} · answered by {t['model'] or 'no model (failed)'} · every number is calculated; see 'How this was worked out'")
                    st.markdown(t["a"])
                    for k, tr in enumerate(t["trace"]):
                        res = tr["result"]
                        if isinstance(res.get("table"), list) and res["table"]:
                            with st.expander(f"Table: {TOOL_NAMES.get(tr['tool'], tr['tool'])} ({len(res['table'])} rows)"):
                                ptab = plain_table(res["table"])
                                st.dataframe(ptab, hide_index=True, width="stretch",
                                             column_config={c: st.column_config.NumberColumn(format="localized") for c in ptab.columns
                                                            if pd.api.types.is_numeric_dtype(ptab[c]) and c != "Item no."})
                                st.download_button("Download as spreadsheet (CSV)", plain_table(res["table"]).to_csv(index=False),
                                                   file_name=f"Q{t['n']}_{tr['tool']}.csv", key=f"dl{t['n']}_{k}")
                        if res.get("chart") and res["chart"].get("data"):
                            cdf = pd.DataFrame(res["chart"]["data"])
                            cdf["label"] = [short(x) if x in docs else x for x in cdf["label"]]
                            title = res["chart"]["title"]
                            if "(INR)" in title:
                                cdf["value"] = cdf["value"] / 1e5; title = title.replace("(INR)", "(Rs lakh)")
                            st.caption(title)
                            st.bar_chart(cdf.set_index("label"), horizontal=True, height=60 + 40 * len(cdf), color="#7f39ff")
                    if t["trace"]:
                        with st.expander("How this was worked out"):
                            for tr in t["trace"]:
                                st.markdown(f"**{TOOL_NAMES.get(tr['tool'], tr['tool'])}**" + (f" ({plain_args(tr['args'])})" if plain_args(tr['args']) else ""))
                                st.write(tr["result"].get("calculation", ""))
                                for cv in tr["result"].get("caveats", []):
                                    for v in docs:
                                        cv = cv.replace(v, short(v))
                                    st.write(f"- {cv}")

# ---------------------------------------------------------------- tab 4
with t4:
    if not cells:
        st.info("Nothing to review yet.")
    else:
        drafts = {v: d for v in docs if (d := followup.draft(v, short(v), docs[v], items, cells, vflags, S.get("rfq") or {}, S.last_year,
                                                               S.get("who") or "Buyer"))}
        if drafts:
            with st.container(border=True):
                section("Ask the vendors instead of guessing",
                        f"Every open question becomes one email per vendor ({sum(d['points'] for d in drafts.values())} points across "
                        f"{len(drafts)} vendors). Drafts update as you decide below: settle a point here and it leaves the email.")
                fv = st.pills("Vendor", list(drafts), key="fu_vendor", default=list(drafts)[0],
                              format_func=lambda v: f"{short(v)} ({drafts[v]['points']})", label_visibility="collapsed")
                d = drafts.get(fv) or next(iter(drafts.values()))
                st.markdown(f"**Subject:** {d['subject']}")
                st.code(d["body"], language=None, wrap_lines=True)
                f1, f2, _ = st.columns([1.3, 1.6, 2.5])
                if f1.button("Mark as sent", key=f"fu_sent_{d['vendor']}", width="stretch",
                             help="Recorded in the decision record. Sending itself is simulated in this prototype."):
                    tot = split_total(S.resolutions, fx)
                    record(d["vendor"], 0, "Follow-up email", f"Marked as sent ({d['points']} points)", d["subject"], tot, tot)
                    st.toast(f"Recorded: follow-up to {short(d['vendor'])}")
                f2.download_button("Download all drafts (.txt)", width="stretch", file_name="follow_up_emails.txt",
                                   data="\n\n" .join(f"To: {short(x['vendor'])}\nSubject: {x['subject']}\n\n{x['body']}\n" + "-" * 60
                                                       for x in drafts.values()))
        fl = all_open_flags(cells, vflags)
        if not fl:
            st.success("Nothing left to check. Every price is confirmed, converted or marked not quoted.")
        else:
            section(f"{len(fl)} things to check", "The app will not guess these for you. Each decision updates the comparison and the AI's "
                                                  "answers, and is written to the decision record (tab 6) with the total before and after.")
            reset = lambda: S.update(flag_page=1)
            counts = {}
            for f in fl:
                counts[f["type"]] = counts.get(f["type"], 0) + 1
            # Options are the stable type codes; only the label carries the count, so a selection survives a decision changing it.
            type_opts = [k for k, _ in sorted(counts.items(), key=lambda kv: -kv[1])]
            pick_types = st.pills("Type", type_opts, selection_mode="multi", key="flag_types", on_change=reset,
                                  format_func=lambda k: f"{FLAG_NAMES.get(k, k)} ({counts.get(k, 0)})",
                                  help="Pick one or more types. None selected shows all.")
            vcounts = {}
            for f in fl:
                vcounts[f["vendor"]] = vcounts.get(f["vendor"], 0) + 1
            vopts = ["All vendors"] + list(vcounts)
            pick_v = st.pills("Vendor", vopts, key="flag_vendor", on_change=reset, default="All vendors",
                              format_func=lambda v: v if v == "All vendors" else f"{short(v)} ({vcounts.get(v, 0)})")
            want_types = set(pick_types or []) & set(type_opts)
            want_v = None if pick_v in (None, "All vendors") else pick_v
            shown = [f for f in fl if (not want_types or f["type"] in want_types) and (not want_v or f["vendor"] == want_v)]
            if not shown:
                st.caption("Nothing to check matches these filters.")
            a, b = pager(len(shown), 10, "flag_page")
            for f in shown[a:b]:
                k = f["key"]
                with st.container(border=True):
                    where = f"item {f['item_id']}. {next((it['desc'] for it in items if it['id'] == f['item_id']), '')}" if f["item_id"] else "whole quote"
                    st.markdown(f"**{FLAG_NAMES.get(f['type'], f['type'])}** · {short(f['vendor'])} · {where}")
                    st.write(f["message"])
                    if f["item_id"]:
                        c = cells[f["vendor"]][f["item_id"]]
                        if c["source"]:
                            st.caption(f"Source {c['source'].get('ref')}: {c['source'].get('snippet')}")
                            show_crop(f["vendor"], c)
                    b1, b2 = st.columns(2)
                    if f["type"] in ("LOW_CONFIDENCE", "UNIT_MISMATCH", "CURRENCY_UNKNOWN", "BASELINE_ASSUMED", "UNREADABLE_PRICE") and f["item_id"]:
                        c = cells[f["vendor"]][f["item_id"]]
                        val = b2.number_input("Or type the correct price (Rs, per unit asked for)", value=(None if c["display_price"] is None else float(c["display_price"])), min_value=0.0, step=0.5,
                                              placeholder="Type the price you read", key=f"v_{k}")
                        if b2.button("Use my value", key=f"e_{k}", disabled=not val):
                            decide(k, dict(action="edit", price=val), f["vendor"], f["item_id"], FLAG_NAMES.get(f["type"], f["type"]),
                                   f"Typed price Rs {val:g}", source_text(f))
                    if f["type"] == "UNMATCHED_LINE":
                        target = b2.selectbox("Which item is it?", [f"{it['id']}. {it['desc']}" for it in items], key=f"m_{k}")
                        if b2.button("Match to this item", key=f"mb_{k}"):
                            decide(k, dict(action="map", item_id=int(target.split(".")[0])), f["vendor"], 0, FLAG_NAMES.get(f["type"], f["type"]),
                                   f"Matched to item {target}", f["message"])
                    if f["type"] != "UNREADABLE_PRICE":
                        label = ("Leave it out" if f["type"] == "UNMATCHED_LINE" else "Accept the claim" if f["type"] == "EVIDENCE_MISSING"
                                 else "Looks right" if f["item_id"] else "Accept")
                        if b1.button(label, key=f"a_{k}"):
                            decide(k, dict(action="accept"), f["vendor"], f["item_id"], FLAG_NAMES.get(f["type"], f["type"]), label, source_text(f))

# ---------------------------------------------------------------- tab 5
with t5:
    if not cells:
        st.info("Nothing decided yet. Load the request and the vendor replies first.")
    else:
        split = split_now
        nopen = len(all_open_flags(cells, vflags))
        section("Who gets the order (split award), as things stand", "Updated with every decision. Download the award pack to attach to the approval email.")
        k1, k2, k3 = st.columns(3)
        k1.metric("Total yearly cost", fmt_inr(split["total_inr"]))
        k2.metric("Saving vs last year", fmt_inr(split["saving_vs_last_year_inr"]), f"{split['saving_vs_last_year_pct']}%")
        k3.metric("Status", "Final" if not nopen else "Draft", help="Draft until every thing to check has a decision.")
        st.caption(" · ".join(f"{short(v)}: {fmt_inr(x)}" for v, x in sorted(split["by_vendor_inr"].items(), key=lambda kv: -kv[1])))
        if nopen:
            st.warning(f"{nopen} things still to check (tab 5). Until they are decided, those prices are left out and the pack is marked DRAFT.")
        pack = award_pack.build(docs, items, cells, vflags, S.last_year, fx, S.log, S.turns, S.get("who") or "Buyer", now_ist(),
                                short=short, flag_names=FLAG_NAMES, status_names=STATUS_NAMES)
        st.download_button("Download award pack (Excel)", pack, type="primary",
                           file_name=f"award_pack_{pd.Timestamp.now(tz='Asia/Kolkata'):%Y%m%d_%H%M}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           help="Summary, who gets each item, the full comparison, quality evidence, this decision record, what is still "
                                "to check, and the questions asked. Attach it to the approval email.")

        section(f"Decision record ({len(S.log)} entries, newest first)", "Who decided what, when, based on what, and how it moved the total. "
                                                                          "Entries are never edited or deleted: an undo is a new entry, like a register.")
        if not S.log:
            st.caption("No decisions yet. Decide things in tab 5 and they appear here.")
        else:
            entries = list(reversed(S.log))
            a, b = pager(len(entries), 10, "log_page")
            def change(e):
                if None in (e["before"], e["after"]):
                    return ""
                return "no change" if e["before"] == e["after"] else f"{'+' if e['after'] > e['before'] else '−'}{fmt_inr(abs(e['after'] - e['before']))}"
            st.dataframe(pd.DataFrame([{"#": e["n"], "When": pd.Timestamp(e["at"].replace(" IST", "")).strftime("%d %b, %H:%M") if e["at"] else "",
                                        "Who": e["who"], "Vendor": e["vendor"], "Item": e["item"],
                                        "Decision": (f"{e['what']}: {e['decision'][:1].lower()}{e['decision'][1:]}" if e["what"] else e["decision"]).replace("typed price Rs", "typed Rs"),
                                        "Change": change(e), "Total after": fmt_inr(e["after"])} for e in entries[a:b]]),
                         hide_index=True, width="stretch",
                         column_config={"#": st.column_config.NumberColumn(width=34), "When": st.column_config.TextColumn(width=98),
                                        "Who": st.column_config.TextColumn(width=62), "Vendor": st.column_config.TextColumn(width=140),
                                        "Item": st.column_config.TextColumn(width=200), "Decision": st.column_config.TextColumn(width=232),
                                        "Change": st.column_config.TextColumn("Change to total", width=110),
                                        "Total after": st.column_config.TextColumn(width=86)})
            with st.expander("What each decision was based on"):
                for e in entries[a:b]:
                    st.markdown(f"**#{e['n']}** · {e['vendor'] or 'Setting'} · {e['item'] or e['what']}  \n{e['source']}")
            active = [e for e in reversed(S.log) if e.get("key") in S.resolutions and not e["decision"].startswith("Undid")]
            seen, choices = set(), []
            for e in active:
                if e["key"] not in seen:
                    seen.add(e["key"]); choices.append(e)
            if choices:
                u1, u2 = st.columns([4, 1], vertical_alignment="bottom")
                pick = u1.selectbox("Undo a decision", choices, format_func=lambda e: f"#{e['n']} · {e['vendor']} · {e['item']} · {e['decision']}",
                                    key="undo_pick")
                if u2.button("Undo", width="stretch"):
                    undo(pick["key"]); st.rerun()

# ---------------------------------------------------------------- record header (top of the page)
with header_slot:
    rfq = S.get("rfq") or {}
    if S.rfx:
        title = rfq.get("title") or "Request for quotes"
        meta = [f"<b>RFQ {rfq['no']}</b>" if rfq.get("no") else "", rfq.get("site", ""), f"<b>{len(S.rfx)}</b> items",
                f"<b>{len(S.docs)}</b> vendor replies read" if S.docs else "no replies read yet",
                f"<b>{len(S.supports)}</b> certificates and test reports" if S.supports else ""]
    else:
        title, meta = "No request loaded yet", ["Start in tab 1: load the sample request, or describe what you need"]
    if split_now:
        nopen = len(all_open_flags(cells, vflags))
        badge = '<span class="badge final">Final</span>' if not nopen else '<span class="badge draft">Draft</span>'
        right = (f'<div class="eyebrow">Who gets the order (split award)</div><div class="big">{fmt_inr(split_now["total_inr"])}{badge}</div>'
                 f'<div class="meta">saves {fmt_inr(split_now["saving_vs_last_year_inr"])} vs last year · '
                 + (f'<b>{nopen}</b> things to check' if nopen else 'nothing left to check') + '</div>')
    else:
        right = '<div class="meta">The comparison appears once replies are read.</div>'
    st.markdown(f'''<div class="rec"><div class="left"><div class="eyebrow">ClearQuote <span style='font-weight:400;text-transform:none;letter-spacing:0'>· every vendor quote, made comparable</span></div>
<div class="title">{title}</div><div class="meta">{" · ".join(m for m in meta if m)}</div></div><div class="right">{right}</div></div>''',
                unsafe_allow_html=True)
