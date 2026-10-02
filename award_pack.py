"""Award pack: one Excel file a buyer attaches to the approval email. Plain code only; every number comes from the same
functions the app and the analyst use, so the file and the screen can never disagree."""
import io
import pandas as pd
from normalize import summarise, all_open_flags, vendor_eligibility, EVIDENCE_FLAGS
from analyst import make_tools


def build(docs, items, cells, vflags, last_year, fx, log, turns, who, generated_at, short=lambda v: v, flag_names=None,
          status_names=None):
    """Return the award pack as .xlsx bytes. docs must already carry '_evidence' (normalize.attach_evidence)."""
    flag_names, status_names = flag_names or {}, status_names or {}
    tools = {f.__name__: f for f in make_tools(docs, items, cells, vflags, last_year, [])}
    split = tools["split_award"](only_eligible=True)
    open_fl = all_open_flags(cells, vflags)
    rows, ready = summarise(docs, items, cells, vflags)
    desc = {it["id"]: it["desc"] for it in items}

    # Summary
    summary = [("Who gets the order (split award)", ""),
               ("Prepared by", who), ("Generated", generated_at), ("Dollar to rupee rate used", fx),
               ("", ""),
               ("Total yearly cost (Rs)", split["total_inr"]),
               ("Last year's cost for the same items (Rs)", split["last_year_cost_same_lines_inr"]),
               ("Saving vs last year (Rs)", split["saving_vs_last_year_inr"]),
               ("Saving vs last year (%)", split["saving_vs_last_year_pct"]),
               ("Extra cost of the cheapest single vendor instead (Rs)", split["best_single_award_saving_inr"]),
               ("Items with no usable price (left out of the total)", ", ".join(map(str, split["lines_with_no_usable_price"])) or "none"),
               ("", "")]
    summary += [(f"Order value: {short(v)} (Rs)", amt) for v, amt in sorted(split["by_vendor_inr"].items(), key=lambda kv: -kv[1])]
    summary += [("", ""),
                ("Comparison complete", f"{ready:.0%}"),
                ("Things still to check", len(open_fl)),
                ("Status", "FINAL: every price has been checked." if not open_fl else
                 f"DRAFT: {len(open_fl)} things still to check. Prices waiting on a decision are left out of every total."),
                ("", "")]
    summary += [("Watch out for", short_text(c, docs, short)) for c in split["caveats"]]
    summary += [("", ""), ("How this was worked out", split["calculation"]),
                ("Note", "AI only read the vendors' documents. Every number in this file was calculated by code, "
                         "and every buyer decision is listed in the 'Decision log' sheet.")]

    award = pd.DataFrame([{"Item no.": r["item_id"], "Item": r["item"], "Yearly quantity": r["qty"], "Vendor": short(r["vendor"]),
                           "Price (Rs)": r["price_inr"], "Cost (Rs)": r["cost_inr"],
                           "How the price was worked out": cells[r["vendor"]][r["item_id"]]["receipt"],
                           "Source": (cells[r["vendor"]][r["item_id"]]["source"] or {}).get("ref", "")} for r in split["table"]])

    comp = []
    for it in items:
        row = {"Item no.": it["id"], "Item": it["desc"], "Unit": it["unit"], "Yearly quantity": it["qty"]}
        for v in docs:
            c = cells[v][it["id"]]
            row[short(v)] = round(c["price"], 2) if c["price"] is not None else status_names.get(c["status"], c["status"])
        comp.append(row)

    quality = pd.DataFrame([{"Vendor": short(v), "ISO 9001 (said)": d["questionnaire"]["iso_9001"], "ISO 9001 evidence": d["_evidence"]["iso"],
                             "Burst report (said)": d["questionnaire"]["burst_report"], "Burst test evidence": d["_evidence"]["burst"],
                             "Quality check": "Pass" if vendor_eligibility(d)[0] else "Fail",
                             "Documents attached": ", ".join(d["_evidence"]["files"]) or "none",
                             "Unproven claims still open": "; ".join(f["message"] for f in vflags[v] if f["type"] in EVIDENCE_FLAGS and not f["resolved"])}
                            for v, d in docs.items()])
    vendors = pd.DataFrame([{"Vendor": short(r["vendor"]), "Quality check": "Pass" if r["eligible"] else f"Fail: {r['eligibility_note']}",
                             "Items quoted": r["lines_quoted"], "Not re-quoted (last year's price)": r["lines_assumed_from_last_year"],
                             "Items with a usable price": r["lines_usable"], "Things still to check": r["open_flags"]} for r in rows])
    decisions = pd.DataFrame([{"#": e["n"], "When": e["at"], "Who": e["who"], "Vendor": e["vendor"], "Item": e["item"], "What": e["what"],
                               "Decision": e["decision"], "Based on": e["source"], "Total before (Rs)": e["before"],
                               "Total after (Rs)": e["after"], "Change (Rs)": (e["after"] - e["before"]) if None not in (e["before"], e["after"]) else None}
                              for e in log], columns=["#", "When", "Who", "Vendor", "Item", "What", "Decision", "Based on",
                                                       "Total before (Rs)", "Total after (Rs)", "Change (Rs)"])
    still = pd.DataFrame([{"Vendor": short(f["vendor"]), "Item": f"{f['item_id']}. {desc.get(f['item_id'], '')}" if f["item_id"] else "whole quote",
                           "What": flag_names.get(f["type"], f["type"]), "Details": f["message"]} for f in open_fl],
                         columns=["Vendor", "Item", "What", "Details"])
    asked = pd.DataFrame([{"#": t["n"], "When": t["at"], "Question": t["q"], "Answer": t["a"], "Answered by": t.get("model") or ""} for t in turns],
                         columns=["#", "When", "Question", "Answer", "Answered by"])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame(summary, columns=["", " "]).to_excel(xw, sheet_name="Summary", index=False, header=False)
        sheets = [("Who gets the order", award), ("Comparison", pd.DataFrame(comp)), ("Vendors", vendors), ("Quality evidence", quality),
                  ("Decision log", decisions), ("Still to check", still), ("AI questions", asked)]
        for name, df in sheets:
            df.to_excel(xw, sheet_name=name, index=False)
        _style(xw.book)
    return buf.getvalue()


def short_text(text, docs, short):
    for v in docs:
        text = text.replace(v, short(v))
    return text


def _style(wb):
    from openpyxl.styles import Font, Alignment, PatternFill
    head = PatternFill("solid", fgColor="EBE8FB")
    for ws in wb.worksheets:
        widths = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    widths[c.column_letter] = min(max(widths.get(c.column_letter, 0), len(str(c.value))), 70)
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                        c.number_format = "#,##0.00" if isinstance(c.value, float) and c.value != int(c.value) else "#,##0"
        for col, w in widths.items():
            ws.column_dimensions[col].width = max(10, w + 2)
        if ws.title == "Summary":
            ws["A1"].font = Font(bold=True, size=14, color="3D05C6")
            for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
                row[0].font = Font(bold=True)
        else:
            for c in ws[1]:
                c.font, c.fill = Font(bold=True), head
            ws.freeze_panes = "A2"
